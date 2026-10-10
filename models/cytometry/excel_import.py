"""Import a cytometry collection workbook as deposit drafts.

One spreadsheet row is one taxon. Rows that share the FCS filename in
„Měření: Název FCS souboru (KLÍČ)“ become one Invenio draft, with those
taxa in ``metadata.sample_provenance.taxon``.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

DATA_SHEET = "Sběr dat Cytometrie"
GUIDE_SHEET = "Návod"

DATE_RE = re.compile(r"^[0-9]{4}(-[0-9]{2}(-[0-9]{2})?)?$")
RESOURCE_TYPE_ID = "datasets"

LICENSE_ALIASES = {
    "cc-by-4.0": "4-BY",
    "cc-by": "4-BY",
    "CC-BY-4.0": "4-BY",
}
LANGUAGE_ALIASES = {
    "eng": "en",
    "en": "en",
    "cze": "cs",
    "ces": "cs",
    "cs": "cs",
}

SECTION_FILLS = {
    "Projekt": "D6EAF8",
    "Vzorek": "D5F5E3",
    "Lokalita": "E8F8F5",
    "Taxon": "FCF3CF",
    "Příprava": "FAD7A0",
    "Standard": "F9E79F",
    "Přístroj": "E5E8E8",
    "Měření": "E8DAEF",
}


@dataclass(frozen=True)
class Column:
    """One workbook column and how its cells are interpreted."""

    key: str
    header: str
    section: str
    kind: str  # text, list, date, number, ratio, file, files


COLUMNS: tuple[Column, ...] = (
    Column("project_title", "Projekt: Název", "Projekt", "text"),
    Column(
        "project_description",
        "Projekt: Popis",
        "Projekt",
        "text",
    ),
    Column(
        "project_authors",
        "Projekt: Autoři (oddělovat středníkem ;)",
        "Projekt",
        "list",
    ),
    Column("publication_doi", "Projekt: DOI", "Projekt", "text"),
    Column("specimen_id", "Vzorek: ID jedince", "Vzorek", "text"),
    Column("collection_id", "Vzorek: ID sběru", "Vzorek", "text"),
    Column(
        "date_of_sampling",
        "Vzorek: Datum sběru (RRRR, RRRR-MM nebo RRRR-MM-DD)",
        "Vzorek",
        "date",
    ),
    Column(
        "collectors",
        "Vzorek: Sběratel(é) (oddělovat středníkem ;)",
        "Vzorek",
        "list",
    ),
    Column("voucher", "Vzorek: Herbář / Voucher", "Vzorek", "text"),
    Column("material_entity", "Vzorek: Tkáň / Část", "Vzorek", "text"),
    Column("country_or_region", "Lokalita: Země / Region", "Lokalita", "text"),
    Column("locality_description", "Lokalita: Popis", "Lokalita", "text"),
    Column("coordinates", "Lokalita: Souřadnice", "Lokalita", "text"),
    Column("scientific_name", "Taxon: Vědecké jméno (POVINNÉ)", "Taxon", "text"),
    Column("strain_cultivar", "Taxon: Klon / Kultivar", "Taxon", "text"),
    Column("col_id", "Taxon: CoL ID", "Taxon", "text"),
    Column("taxon_notes", "Taxon: Poznámky", "Taxon", "text"),
    Column("isolation_buffer", "Příprava: Izolační pufr", "Příprava", "text"),
    Column("buffer_additives", "Příprava: Aditiva", "Příprava", "text"),
    Column("fluorochrome", "Příprava: Fluorochrom", "Příprava", "text"),
    Column("internal_standard_species", "Standard: Druh", "Standard", "text"),
    Column("internal_standard_material", "Standard: Tkáň", "Standard", "text"),
    Column(
        "standard_c_value_pg",
        "Standard: C-value (pg, číslo)",
        "Standard",
        "number",
    ),
    Column("standard_c_value_citation", "Standard: Citace C-value", "Standard", "text"),
    Column("manufacturer", "Přístroj: Výrobce", "Přístroj", "text"),
    Column("instrument_model", "Přístroj: Model", "Přístroj", "text"),
    Column("instrument_id", "Přístroj: Sériové číslo", "Přístroj", "text"),
    Column("light_source", "Přístroj: Zdroj světla", "Přístroj", "text"),
    Column(
        "measurer_name",
        "Měření: Operátor(ři) (oddělovat středníkem ;)",
        "Měření",
        "list",
    ),
    Column(
        "measurement_date",
        "Měření: Datum měření (RRRR, RRRR-MM nebo RRRR-MM-DD)",
        "Měření",
        "date",
    ),
    Column("purpose", "Měření: Účel", "Měření", "text"),
    Column(
        "peak_ratio",
        "Měření: Poměr píků vzorek/standard",
        "Měření",
        "ratio",
    ),
    Column("fcs_file", "Měření: Název FCS souboru (KLÍČ)", "Měření", "file"),
    Column(
        "extra_files",
        "Měření: Další soubory (oddělovat středníkem ;)",
        "Měření",
        "files",
    ),
    Column("interpretation_notes", "Měření: Poznámky / CV", "Měření", "text"),
)

COLUMNS_BY_KEY = {column.key: column for column in COLUMNS}
TAXON_KEYS = ("scientific_name", "strain_cultivar", "col_id", "taxon_notes")
SHARED_KEYS = tuple(column.key for column in COLUMNS if column.key not in TAXON_KEYS)

GUIDE_LINES = (
    "Návod k sběrné tabulce cytometrie",
    "",
    "Jeden řádek = jeden taxon. Řádky se stejným textem ve sloupci „Měření: Název FCS souboru (KLÍČ)“ jsou jedno měření a vznikne z nich jeden záznam.",
    "Ukázka na listu Sběr dat Cytometrie: Arabidopsis thaliana a Cardamine amara sdílejí fcs_sample_001.fcs, takže import složí jeden záznam se dvěma taxony. Solanum tuberosum má vlastní soubor a vlastní záznam.",
    "Ukázkové řádky před ostrým importem smažte.",
    "",
    "Ostatní sloupce ve skupině musí mít stejnou hodnotu. Prázdná buňka se doplní z ostatních řádků skupiny. Dvě různé neprázdné hodnoty jsou chyba a celá skupina se odmítne.",
    "Seznamy (autoři, sběratelé, operátoři, další soubory) oddělujte středníkem.",
    "Datum pište jako RRRR, RRRR-MM nebo RRRR-MM-DD.",
    "C-value je číslo v pikogramech bez jednotky (2.50) a musí být větší než 0.",
    "Poměr píků vzorek/standard je číslo větší než 0 a nejvýše 500.",
    "Vědecké jméno je na každém řádku povinné.",
    "Název FCS souboru je povinný. Musí přesně odpovídat souboru ve složce, kterou kurátor předá při importu. Cesta (složka/soubor) se do buňky nepíše.",
    "Další soubory jsou volitelné přílohy téhož záznamu, také oddělené středníkem.",
    "",
    "Název záznamu, jazyk, licence a komunita se do tabulky nepíšou. Doplní je import.",
    "Název se složí z vědeckých jmen, ID jedince a data měření.",
    "Autoři projektu se zapíšou i jako tvůrci záznamu. Jméno ve tvaru „Jan Novák“, nebo „Novák, Jan“.",
    "Bez autorů projektu záznam založit nejde.",
)


@dataclass
class SheetRow:
    """One data row after cell parsing. ``values`` holds normalized Python values."""

    number: int
    values: dict[str, Any]
    errors: list[str] = field(default_factory=list)


@dataclass
class MeasurementGroup:
    """Rows that share one FCS key, ready to become one draft."""

    fcs_key: str
    row_numbers: list[int]
    payload: dict[str, Any] | None = None
    files: list[Path] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def header_for(key: str) -> str:
    """Czech column title used in error messages."""
    return COLUMNS_BY_KEY[key].header


def split_list(value: str) -> list[str]:
    """Split a semicolon-separated cell into stripped non-empty items."""
    return [part.strip() for part in value.split(";") if part.strip()]


def is_valid_date(value: str) -> bool:
    """Return whether ``value`` matches the schema date pattern and calendar."""
    if not DATE_RE.fullmatch(value):
        return False
    parts = value.split("-")
    if len(parts) >= 2 and not 1 <= int(parts[1]) <= 12:
        return False
    if len(parts) == 3:
        year, month, day = (int(part) for part in parts)
        try:
            date(year, month, day)
        except ValueError:
            return False
    return True


def cell_to_text(value: Any) -> str:
    """Normalize a worksheet cell to a stripped string."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def parse_date_cell(value: Any) -> str:
    """Return an EDTF date string, or raise ValueError."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if float(value).is_integer() and 1000 <= int(value) <= 9999:
            return str(int(value))
        raise ValueError("očekává se datum RRRR, RRRR-MM nebo RRRR-MM-DD")
    text = cell_to_text(value)
    if text and not is_valid_date(text):
        raise ValueError("očekává se datum RRRR, RRRR-MM nebo RRRR-MM-DD")
    return text


def parse_number_cell(value: Any, *, maximum: float | None) -> float | None:
    """Parse a positive number. ``maximum`` limits the peak ratio."""
    if value is None or value == "":
        return None
    if isinstance(value, bool) or isinstance(value, (datetime, date)):
        raise ValueError("očekává se číslo")
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = cell_to_text(value).replace(",", ".")
        if not text:
            return None
        if any(character.isalpha() for character in text):
            raise ValueError("očekává se číslo bez jednotky")
        try:
            number = float(text)
        except ValueError as exc:
            raise ValueError("očekává se číslo") from exc
    if number <= 0:
        raise ValueError("musí být větší než 0")
    if maximum is not None and number > maximum:
        raise ValueError(f"musí být nejvýše {maximum:g}")
    return number


def parse_cell(column: Column, value: Any) -> Any:
    """Parse one cell according to the column kind."""
    if column.kind == "date":
        return parse_date_cell(value)
    if column.kind == "number":
        return parse_number_cell(value, maximum=None)
    if column.kind == "ratio":
        return parse_number_cell(value, maximum=500)
    if column.kind == "list":
        return split_list(cell_to_text(value))
    if column.kind == "files":
        return split_list(cell_to_text(value))
    return cell_to_text(value)


def parse_sheet_rows(workbook_path: Path) -> tuple[list[SheetRow], list[str], list[str]]:
    """Read the data sheet.

    Returns rows, blocking errors, and warnings (unknown columns).
    """
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, data_only=False)
    if DATA_SHEET in workbook.sheetnames:
        sheet = workbook[DATA_SHEET]
    else:
        sheet = workbook.active
    header_cells = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    headers = [cell_to_text(value) for value in header_cells]
    index_by_header = {header: index for index, header in enumerate(headers) if header}
    missing = [column.header for column in COLUMNS if column.header not in index_by_header]
    known = {column.header for column in COLUMNS}
    unknown = [header for header in headers if header and header not in known]
    workbook_errors: list[str] = []
    warnings: list[str] = []
    if missing:
        workbook_errors.append(
            "V tabulce chybí sloupce: " + "; ".join(missing)
        )
    if unknown:
        warnings.append(
            "Neznámé sloupce (import je ignoruje): " + "; ".join(unknown)
        )
    if missing:
        return [], workbook_errors, warnings

    rows: list[SheetRow] = []
    for excel_row in sheet.iter_rows(min_row=2):
        raw_cells = [cell.value for cell in excel_row]
        if not any(cell_to_text(value) for value in raw_cells):
            continue
        number = excel_row[0].row
        values: dict[str, Any] = {}
        errors: list[str] = []
        for column in COLUMNS:
            index = index_by_header[column.header]
            cell_value = raw_cells[index] if index < len(raw_cells) else None
            try:
                values[column.key] = parse_cell(column, cell_value)
            except ValueError as exc:
                values[column.key] = None if column.kind in {"number", "ratio"} else ""
                errors.append(f"řádek {number}, {column.header}: {exc}")
        if not values.get("scientific_name"):
            errors.append(
                f"řádek {number}, {header_for('scientific_name')}: vědecké jméno je povinné"
            )
        if not values.get("fcs_file"):
            errors.append(
                f"řádek {number}, {header_for('fcs_file')}: název FCS souboru je povinný"
            )
        rows.append(SheetRow(number=number, values=values, errors=errors))
    if not rows and not workbook_errors:
        workbook_errors.append("Tabulka neobsahuje žádný řádek s daty.")
    return rows, workbook_errors, warnings


def _same(left: Any, right: Any) -> bool:
    return left == right


def merge_shared(rows: list[SheetRow], key: str, errors: list[str], fcs_key: str) -> Any:
    """Take the single non-empty value of a shared column, or report a clash."""
    present: list[tuple[int, Any]] = []
    for row in rows:
        value = row.values.get(key)
        if value in (None, "", []):
            continue
        if any(_same(value, previous) for _, previous in present):
            continue
        present.append((row.number, value))
    if len(present) > 1:
        listed = ", ".join(f"řádek {number}" for number, _ in present)
        errors.append(
            f"skupina {fcs_key}, {header_for(key)}: {listed} mají různou hodnotu"
        )
        return present[0][1]
    if not present:
        kind = COLUMNS_BY_KEY[key].kind
        if kind in {"list", "files"}:
            return []
        return None
    return present[0][1]


def check_filename(name: str, files_dir: Path, label: str, row_numbers: list[int]) -> tuple[Path | None, str | None]:
    """Resolve one attachment name inside ``files_dir``."""
    if "/" in name or "\\" in name or name in {".", ".."} or Path(name).name != name:
        rows = ", ".join(str(number) for number in row_numbers)
        return None, (
            f"řádky {rows}, {label}: „{name}“ musí být jen název souboru, ne cesta"
        )
    path = (files_dir / name).resolve()
    try:
        path.relative_to(files_dir.resolve())
    except ValueError:
        return None, f"{label}: „{name}“ je mimo složku souborů"
    if not path.is_file():
        return None, f"{label}: soubor „{name}“ ve složce {files_dir} chybí"
    return path, None


def person_from_name(full_name: str) -> dict[str, Any]:
    """Map a free-text name onto an RDM personal creator."""
    if "," in full_name:
        family_name, given_name = (part.strip() for part in full_name.split(",", 1))
    else:
        parts = full_name.split()
        if len(parts) == 1:
            family_name, given_name = parts[0], ""
        else:
            given_name = " ".join(parts[:-1])
            family_name = parts[-1]
    person: dict[str, Any] = {"type": "personal", "family_name": family_name}
    if given_name:
        person["given_name"] = given_name
    return {"person_or_org": person}


def build_title(names: list[str], specimen_id: str | None, measurement_date: str | None) -> str:
    """Compose the record title from the taxon names and sample identity."""
    head = " + ".join(names) if names else "Cytometry measurement"
    details = [value for value in (specimen_id, measurement_date) if value]
    if details:
        return f"{head} ({', '.join(details)})"
    return head


def compact(value: Any) -> Any:
    """Drop empty strings, lists and objects so the draft matches a filled form."""
    if isinstance(value, dict):
        cleaned = {}
        for key, item in value.items():
            compacted = compact(item)
            if compacted in (None, "", [], {}):
                continue
            cleaned[key] = compacted
        return cleaned
    if isinstance(value, list):
        return [
            compact(item)
            for item in value
            if compact(item) not in (None, "", [], {})
        ]
    return value


def build_groups(
    rows: list[SheetRow],
    *,
    files_dir: Path,
    license_id: str,
    language_id: str,
    community: str | None,
) -> list[MeasurementGroup]:
    """Group parsed rows by FCS filename and build draft payloads."""
    order: OrderedDict[str, list[SheetRow]] = OrderedDict()
    unkeyed: list[SheetRow] = []
    for row in rows:
        fcs_key = row.values.get("fcs_file") or ""
        if not fcs_key:
            unkeyed.append(row)
            continue
        order.setdefault(fcs_key, []).append(row)

    groups: list[MeasurementGroup] = []
    for row in unkeyed:
        groups.append(
            MeasurementGroup(
                fcs_key="",
                row_numbers=[row.number],
                errors=list(row.errors),
            )
        )

    for fcs_key, grouped_rows in order.items():
        errors = [error for row in grouped_rows for error in row.errors]
        shared = {
            key: merge_shared(grouped_rows, key, errors, fcs_key) for key in SHARED_KEYS
        }
        taxa = []
        names: list[str] = []
        for row in grouped_rows:
            scientific_name = row.values.get("scientific_name") or ""
            if scientific_name:
                names.append(scientific_name)
            taxon = {
                "scientific_name": scientific_name,
                "strain_cultivar": row.values.get("strain_cultivar") or None,
                "col_id": row.values.get("col_id") or None,
                "notes": row.values.get("taxon_notes") or None,
            }
            taxa.append(compact(taxon))

        authors = list(shared.get("project_authors") or [])
        if not authors:
            errors.append(
                f"skupina {fcs_key}, {header_for('project_authors')}: "
                "bez autorů projektu nejde vyplnit tvůrce záznamu"
            )

        file_names = [fcs_key, *list(shared.get("extra_files") or [])]
        unique_names: list[str] = []
        for name in file_names:
            if name and name not in unique_names:
                unique_names.append(name)
        files: list[Path] = []
        row_numbers = [row.number for row in grouped_rows]
        for name in unique_names:
            label = header_for("fcs_file") if name == fcs_key else header_for("extra_files")
            path, error = check_filename(name, files_dir, label, row_numbers)
            if error:
                errors.append(f"skupina {fcs_key}: {error}")
            elif path is not None:
                files.append(path)

        measurement_date = shared.get("measurement_date")
        publication_date = measurement_date or date.today().isoformat()
        metadata: dict[str, Any] = {
            "title": build_title(names, shared.get("specimen_id"), measurement_date),
            "resource_type": {"id": RESOURCE_TYPE_ID},
            "publication_date": publication_date,
            "creators": [person_from_name(author) for author in authors],
            "languages": [{"id": language_id}],
            "rights": [{"id": license_id}],
            "project_info": {
                "project_title": shared.get("project_title"),
                "project_description": shared.get("project_description"),
                "authors": authors,
                "publication_doi": shared.get("publication_doi"),
            },
            "sample_provenance": {
                "taxon": [taxon for taxon in taxa if taxon],
                "location": {
                    "country_or_region": shared.get("country_or_region"),
                    "locality_description": shared.get("locality_description"),
                    "coordinates": shared.get("coordinates"),
                },
                "date_of_sampling": shared.get("date_of_sampling"),
                "collector": shared.get("collectors") or [],
                "collection_id": shared.get("collection_id"),
                "voucher": shared.get("voucher"),
                "material_entity": shared.get("material_entity"),
                "specimen_id": shared.get("specimen_id"),
            },
            "preparation_and_standard": {
                "isolation_buffer": shared.get("isolation_buffer"),
                "buffer_additives": shared.get("buffer_additives"),
                "fluorochrome": shared.get("fluorochrome"),
                "internal_standard_species": shared.get("internal_standard_species"),
                "internal_standard_material": shared.get("internal_standard_material"),
                "standard_c_value_pg": shared.get("standard_c_value_pg"),
                "standard_c_value_citation": shared.get("standard_c_value_citation"),
            },
            "instrumentation": {
                "manufacturer": shared.get("manufacturer"),
                "model": shared.get("instrument_model"),
                "instrument_id": shared.get("instrument_id"),
                "light_source": shared.get("light_source"),
            },
            "measurement": {
                "measurer_name": shared.get("measurer_name") or [],
                "measurement_date": measurement_date,
                "purpose": shared.get("purpose"),
                "peakSampleToPeakStandard": shared.get("peak_ratio"),
                "interpretation_notes": shared.get("interpretation_notes"),
            },
        }
        payload: dict[str, Any] = {
            "metadata": compact(metadata),
            "files": {"enabled": True},
        }
        if community:
            payload["parent"] = {"communities": {"default": {"id": community}}}
        groups.append(
            MeasurementGroup(
                fcs_key=fcs_key,
                row_numbers=row_numbers,
                payload=payload,
                files=files,
                errors=errors,
            )
        )
    return groups


def resolve_license(value: str) -> str:
    """Map a CLI license alias onto the rights vocabulary id."""
    return LICENSE_ALIASES.get(value, value)


def resolve_language(value: str) -> str:
    """Map a CLI language alias onto the languages vocabulary id."""
    return LANGUAGE_ALIASES.get(value, value)


def format_rows(numbers: list[int]) -> str:
    """Format Excel row numbers for a report line."""
    if not numbers:
        return ""
    if numbers == list(range(numbers[0], numbers[-1] + 1)) and len(numbers) > 2:
        return f"řádky {numbers[0]}–{numbers[-1]}"
    if len(numbers) == 1:
        return f"řádek {numbers[0]}"
    return "řádky " + ", ".join(str(number) for number in numbers)


def report(
    groups: list[MeasurementGroup],
    workbook_errors: list[str],
    warnings: list[str] | None = None,
) -> str:
    """Human-readable dry-run report."""
    lines: list[str] = []
    if warnings:
        lines.append("Upozornění:")
        lines.extend(f"- {warning}" for warning in warnings)
        lines.append("")
    errors = list(workbook_errors)
    for group in groups:
        errors.extend(group.errors)
    if errors:
        lines.append("Chyby:")
        lines.extend(f"- {error}" for error in errors)
        lines.append("")
    for group in groups:
        if group.errors or not group.payload:
            continue
        title = group.fcs_key or "(bez FCS)"
        lines.append(f"{title} ({format_rows(group.row_numbers)})")
        file_names = ", ".join(path.name for path in group.files) or "(žádné)"
        lines.append(f"Soubory: {file_names}")
        lines.append(json.dumps(group.payload, ensure_ascii=False, indent=2))
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def load_measurement_groups(
    workbook_path: Path,
    files_dir: Path,
    *,
    license_id: str,
    language_id: str,
    community: str | None,
) -> tuple[list[MeasurementGroup], list[str], list[str]]:
    """Parse a workbook into measurement groups."""
    rows, workbook_errors, warnings = parse_sheet_rows(workbook_path)
    if any(error.startswith("V tabulce chybí sloupce") for error in workbook_errors):
        return [], workbook_errors, warnings
    groups = build_groups(
        rows,
        files_dir=files_dir,
        license_id=license_id,
        language_id=language_id,
        community=community,
    )
    return groups, workbook_errors, warnings


def identity_for_user(identifier: str):
    """Build an identity that can create a draft as the given user."""
    from invenio_access.permissions import any_user, authenticated_user
    from invenio_access.utils import get_identity
    from invenio_accounts.proxies import current_datastore

    user = current_datastore.get_user(identifier)
    if user is None:
        user = current_datastore.find_user(username=identifier)
    if user is None:
        user = current_datastore.find_user(email=identifier)
    if user is None:
        raise SystemExit(f"Uživatel „{identifier}“ v Inveniu neexistuje.")
    identity = get_identity(user)
    identity.provides.add(authenticated_user)
    identity.provides.add(any_user)
    return identity


def create_drafts(groups: list[MeasurementGroup], user: str) -> list[str]:
    """Create one draft per valid group and attach its files. Returns draft ids."""
    from invenio_app.factory import create_app
    from oarepo_runtime.proxies import current_runtime

    app = create_app()
    draft_ids: list[str] = []
    with app.app_context():
        identity = identity_for_user(user)
        model = current_runtime.models["cytometry"]
        service = model.service
        files_service = model.draft_file_service
        if files_service is None:
            raise SystemExit("Model cytometry nemá službu pro soubory draftu.")
        for group in groups:
            if group.payload is None:
                continue
            try:
                item = service.create(identity, group.payload)
            except Exception as exc:
                raise SystemExit(
                    f"skupina {group.fcs_key}: záznam se nepodařilo založit ({exc})"
                ) from exc
            draft_id = item.id
            if group.files:
                files_service.init_files(
                    identity,
                    draft_id,
                    [{"key": path.name} for path in group.files],
                )
                for path in group.files:
                    with path.open("rb") as stream:
                        files_service.set_file_content(
                            identity,
                            draft_id,
                            path.name,
                            stream,
                            path.stat().st_size,
                        )
                    files_service.commit_file(identity, draft_id, path.name)
            draft_ids.append(draft_id)
            print(f"{group.fcs_key}: draft {draft_id}  /cytometry/uploads/{draft_id}")
    return draft_ids


def write_template(path: Path) -> None:
    """Write the collection workbook with example rows and a guide sheet."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = DATA_SHEET

    header_font = Font(bold=True)
    header_alignment = Alignment(wrap_text=True, vertical="center")
    for index, column in enumerate(COLUMNS, start=1):
        cell = sheet.cell(1, index, column.header)
        cell.font = header_font
        cell.alignment = header_alignment
        cell.fill = PatternFill(
            "solid",
            fgColor=SECTION_FILLS[column.section],
        )
        width = min(42, max(18, len(column.header) * 0.7))
        sheet.column_dimensions[get_column_letter(index)].width = width

    comments = {
        "standard_c_value_pg": "Jen číslo, například 2.50. Jednotku pg nepište.",
        "standard_c_value_citation": "Například Doležel et al. 2007.",
        "peak_ratio": "Poměr polohy píku vzorku a standardu. Číslo > 0 a ≤ 500.",
        "extra_files": "Volitelné další přílohy, názvy oddělené středníkem.",
        "measurer_name": "Více operátorů oddělte středníkem.",
        "fcs_file": "Klíč záznamu. Stejný název na více řádcích = jeden záznam s více taxony.",
    }
    key_to_index = {column.key: index for index, column in enumerate(COLUMNS, start=1)}
    for key, text in comments.items():
        sheet.cell(1, key_to_index[key]).comment = Comment(text, "cytometry")

    shared = {
        "project_title": "Celoplošný screening flóry ČR",
        "project_description": "Pilotní projekt mapování genomu",
        "project_authors": "Jan Novák; Petr Svoboda",
        "publication_doi": "10.1111/j.1234.2026.x",
        "specimen_id": "SPEC-2026-001",
        "collection_id": "COL-9982",
        "date_of_sampling": "2026-05-12",
        "collectors": "Jan Novák; Alena Bílá",
        "voucher": "PRC 456123",
        "material_entity": "list",
        "country_or_region": "Česká republika",
        "locality_description": "Krkonoše, louka u potoka",
        "coordinates": "50.735, 15.739",
        "isolation_buffer": "Otto buffer",
        "buffer_additives": "PVP; RNase",
        "fluorochrome": "DAPI",
        "internal_standard_species": "Glycine max",
        "internal_standard_material": "mladý list",
        "standard_c_value_pg": 2.50,
        "standard_c_value_citation": "Doležel et al. 2007",
        "manufacturer": "Sysmex",
        "instrument_model": "CyFlow Space",
        "instrument_id": "SN-987654",
        "light_source": "UV LED 365 nm",
        "measurer_name": "Petr Svoboda",
        "measurement_date": "2026-05-14",
        "purpose": "Ploidy determination",
        "peak_ratio": 1.15,
        "fcs_file": "fcs_sample_001.fcs",
        "extra_files": "",
        "interpretation_notes": "CV < 3%, čistý histogram",
    }
    examples = [
        {
            **shared,
            "scientific_name": "Arabidopsis thaliana",
            "strain_cultivar": "Col-0",
            "col_id": "CoL-12345",
            "taxon_notes": "Dominantní druh na lokalitě",
        },
        {
            **shared,
            "scientific_name": "Cardamine amara",
            "strain_cultivar": "Wildtype",
            "col_id": "CoL-67890",
            "taxon_notes": "Přimíchaná minoritní rostlina",
        },
        {
            "project_title": shared["project_title"],
            "project_description": shared["project_description"],
            "project_authors": shared["project_authors"],
            "publication_doi": shared["publication_doi"],
            "specimen_id": "SPEC-2026-002",
            "collection_id": "COL-9983",
            "date_of_sampling": "2026-05-13",
            "collectors": "Alena Bílá",
            "voucher": "PRC 456124",
            "material_entity": "list",
            "country_or_region": "Česká republika",
            "locality_description": "Botanická zahrada UK",
            "coordinates": "50.072, 14.421",
            "scientific_name": "Solanum tuberosum",
            "strain_cultivar": "Ditta",
            "col_id": "CoL-11223",
            "taxon_notes": "",
            "isolation_buffer": "Sgorbati buffer",
            "buffer_additives": "beta-mercaptoethanol",
            "fluorochrome": "PI",
            "internal_standard_species": "Pisum sativum",
            "internal_standard_material": "děložní lístek",
            "standard_c_value_pg": 9.09,
            "standard_c_value_citation": "Doležel et al. 2007",
            "manufacturer": "BD Biosciences",
            "instrument_model": "FACSCanto II",
            "instrument_id": "SN-334455",
            "light_source": "488 nm laser",
            "measurer_name": "Alena Bílá",
            "measurement_date": "2026-05-15",
            "purpose": "Genome size estimation",
            "peak_ratio": 0.42,
            "fcs_file": "fcs_sample_002.fcs",
            "extra_files": "",
            "interpretation_notes": "Standardní měření",
        },
    ]
    for row_index, example in enumerate(examples, start=2):
        for column_index, column in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row_index, column_index, example[column.key])
            if column.kind in {"number", "ratio"}:
                cell.number_format = "0.00"

    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 45
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{1 + len(examples)}"
    sheet.sheet_properties.tabColor = "1A5276"

    guide = workbook.create_sheet(GUIDE_SHEET)
    guide.sheet_properties.tabColor = "196F3D"
    guide.column_dimensions["A"].width = 120
    for row_index, line in enumerate(GUIDE_LINES, start=1):
        cell = guide.cell(row_index, 1, line)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        if row_index == 1:
            cell.font = Font(bold=True, size=14)
    guide.row_dimensions[1].height = 24
    for row_index in range(2, len(GUIDE_LINES) + 1):
        guide.row_dimensions[row_index].height = 32

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    """Command line of the importer."""
    parser = argparse.ArgumentParser(
        description="Naimportuje sběrnou tabulku cytometrie jako drafty záznamů.",
    )
    parser.add_argument(
        "workbook",
        nargs="?",
        type=Path,
        help="Vyplněný soubor .xlsx",
    )
    parser.add_argument(
        "--files-dir",
        type=Path,
        help="Složka s FCS a dalšími přílohami. Názvy musí odpovídat tabulce.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Jen sestaví záznamy a vypíše chyby. Nic nezakládá.",
    )
    parser.add_argument(
        "--user",
        help="E-mail nebo uživatelské jméno vlastníka draftů.",
    )
    parser.add_argument(
        "--license",
        default="cc-by-4.0",
        help="Id licence ve slovníku práv, nebo alias cc-by-4.0 (= 4-BY).",
    )
    parser.add_argument(
        "--language",
        default="eng",
        help="Id jazyka ve slovníku, nebo alias eng (= en).",
    )
    parser.add_argument(
        "--community",
        help="Slug komunity. Bez něj vznikne individuální draft.",
    )
    parser.add_argument(
        "--write-template",
        type=Path,
        metavar="PATH",
        help="Zapíše prázdnou šablonu s ukázkou a skončí.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    """Run a dry-run or create drafts. Returns a process status code."""
    args = parse_args(argv)
    if args.write_template:
        write_template(args.write_template)
        print(f"Šablona zapsána: {args.write_template}")
        return 0
    if args.workbook is None or args.files_dir is None:
        print("Zadejte workbook a --files-dir, nebo --write-template.", file=sys.stderr)
        return 2
    if not args.workbook.is_file():
        print(f"Soubor {args.workbook} neexistuje.", file=sys.stderr)
        return 2
    if not args.files_dir.is_dir():
        print(f"Složka {args.files_dir} neexistuje.", file=sys.stderr)
        return 2
    if not args.dry_run and not args.user:
        print("Ostrý import potřebuje --user.", file=sys.stderr)
        return 2

    license_id = resolve_license(args.license)
    language_id = resolve_language(args.language)
    groups, workbook_errors, warnings = load_measurement_groups(
        args.workbook,
        args.files_dir,
        license_id=license_id,
        language_id=language_id,
        community=args.community,
    )
    print(
        f"Nastavení: licence {license_id}, jazyk {language_id}, "
        f"komunita {args.community or '(žádná)'}"
    )
    print(report(groups, workbook_errors, warnings), end="")
    failed = bool(workbook_errors) or any(group.errors for group in groups)
    if failed:
        return 1
    if args.dry_run:
        return 0
    create_drafts(groups, args.user)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
